import csv, json, collections, pathlib, re
import openpyxl

root=pathlib.Path('outputs/caris-raw-data/TN26-279853')
out=pathlib.Path('work/oct1-analysis')
records=[]; meta=[]
for line_no,line in enumerate((root/'DNA_TN26-279853.vcf').open(),1):
    line=line.rstrip('\n')
    if line.startswith('##'): meta.append(line); continue
    if line.startswith('#'): header=line.lstrip('#').split('\t'); continue
    a=line.split('\t'); info={}
    for item in a[7].split(';'):
        x=item.split('=',1); info[x[0]]=x[1] if len(x)>1 else True
    r=dict(zip(header,a)); r['line']=line_no;r['info']=info
    r['sample_format']=dict(zip(a[8].split(':'),a[9].split(':')))
    records.append(r)

w=openpyxl.load_workbook(root/'TN26-279853_20260925225733.xlsx',read_only=True,data_only=True)
s=w.active;s.reset_dimensions(); rows=list(s.iter_rows(values_only=True));h=rows[0]
seen=collections.Counter(); unique_h=[]
for k in h:
    seen[k]+=1;unique_h.append(k if seen[k]==1 else k+' ('+str(seen[k])+')')
keep=['Accession Number','Case Collection Date','Case Signed Date','Amendment Comments','Amendment Released Date','Client Specimen Id','Specimen Site','Specimen Type','Primary Tumor Site','Lineage','Sub Lineage','Test Created Time','Test Signed Time']
data=[{'source_row':i+2,**{k:v for k,v in zip(unique_h,row) if k in keep or unique_h.index(k)>=53}} for i,row in enumerate(rows[1:])]
res={
 'vcf_header':header,'vcf_metadata':meta,'vcf_record_count':len(records),
 'vcf_filters':dict(collections.Counter(r['FILTER'] for r in records)),
 'vcf_clinical_impact':dict(collections.Counter(str(r['info'].get('CI')) for r in records)),
 'vcf_functional_consequences':dict(collections.Counter(str(r['info'].get('FC')) for r in records)),
 'xlsx_sheet':s.title,'xlsx_actual_rows_including_header':len(rows),'xlsx_actual_columns':len(h),
 'xlsx_headers':[{ 'column_index':i+1,'column_letter':openpyxl.utils.get_column_letter(i+1),'header':k} for i,k in enumerate(unique_h)],
 'xlsx_tests':dict(collections.Counter(r['Test'] for r in data)),
 'xlsx_results':dict(collections.Counter(r['Test Result'] for r in data)),
 'xlsx_technologies':dict(collections.Counter(r['Technology'] for r in data)),
 'xlsx_dimension_bug':'Declared A1:A1; reset_dimensions required. Actual 1277 rows x 128 columns.',
}
for name,obj in [('variants.inventory.json',res),('variants.vcf-records.json',records),('variants.xlsx-records.json',data)]:
    (out/name).write_text(json.dumps(obj,indent=2,default=str))
with (out/'variants.vcf-extracted.tsv').open('w') as f:
    fields=['line','CHROM','POS','REF','ALT','QUAL','FILTER','GI','TI','PC','DC','FC','CI','DP','dbSNP','GT','VF','AD','SB','SA']
    cw=csv.DictWriter(f,fields,delimiter='\t');cw.writeheader()
    for r in records: cw.writerow({k:r.get(k,r['info'].get(k,r['sample_format'].get(k,''))) for k in fields})
print(json.dumps(res,indent=2))
print('\nNONNEGATIVE XLSX RESULTS')
print(json.dumps([r for r in data if r['Test Result'] not in ('variantnotdetected','Not Detected','Negative','Normal')],indent=2))
